import os
import json
import tempfile
import threading
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from PySide6.QtCore import QDate, QRectF, QSize
from PySide6.QtGui import QPainter, QPdfWriter
from PySide6.QtWidgets import QApplication, QDialog, QMessageBox
from docx import Document

from app import (
    ActPreviewDialog, MainWindow, is_excluded_executor, read_lines, write_lines,
    canonical_equipment_model, executor_name_with_initials,
    normalize_service_time_input,
    service_datetime_for_doc,
    validate_service_period,
)
from version import __version__
from app_paths import TEMPLATE_TYPE1, TEMPLATE_TYPE2


class UiSmokeTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.application = QApplication.instance() or QApplication([])

    def test_main_window_builds_all_primary_controls(self):
        window = MainWindow()
        try:
            self.assertEqual("Генератор актов", window.windowTitle())
            self.assertIsNotNone(window.btn_create)
            self.assertIsNotNone(window.cb_executor)
            self.assertIsNotNone(window.cb_location)
            self.assertIsNotNone(window.cb_model)
            self.assertIsNotNone(window.act_select_excel)
            self.assertIsNotNone(window.act_use_embedded_excel)
            self.assertIsNotNone(window.act_open_excel)
            self.assertIsNotNone(window.act_update_equipment_registry)
            self.assertTrue(window.act_preview_enabled.isCheckable())
            self.assertTrue(window.act_batch_panel.isCheckable())
            self.assertTrue(window.act_history_panel.isCheckable())
            self.assertFalse(hasattr(window, "btn_modes"))
            file_actions = window.menu_file.actions()
            self.assertIn(window.act_preview_enabled, file_actions)
            self.assertIn(window.act_batch_panel, file_actions)
            self.assertIn(window.act_history_panel, file_actions)
            self.assertIn(window.act_update_equipment_registry, file_actions)
            self.assertTrue(window.batch_panel.isHidden())
            self.assertTrue(window.history_panel.isHidden())
            self.assertTrue(window.ed_date.calendarPopup())
            self.assertFalse(window.ed_date.isReadOnly())
            self.assertFalse(window.service_period_widget.isHidden())
        finally:
            window.close()

    def test_service_period_is_shown_only_for_mid_validator(self):
        window = MainWindow()
        try:
            window.cb_template.setCurrentIndex(0)
            window.cb_model.setEditText("МИД")
            self.application.processEvents()
            self.assertFalse(window.service_period_widget.isHidden())
            self.assertEqual(
                "Validator_MID.docx",
                Path(window.selected_template(window.cb_model.text())).name,
            )
            window.show()
            self.application.processEvents()
            self.assertGreater(
                window.service_period_widget.geometry().top(),
                window.cb_template.geometry().top(),
            )
            window.cb_model.setEditText("МКТФ")
            self.application.processEvents()
            self.assertTrue(window.service_period_widget.isHidden())
        finally:
            window.close()

    def test_service_period_format_and_order(self):
        self.assertEqual(
            "08:30 часов «17» августа 2026 года",
            service_datetime_for_doc("08:30", "17.08.2026"),
        )
        with self.assertRaises(ValueError):
            validate_service_period("10:00", "09:00", "17.08.2026")
        self.assertEqual(
            "08:30",
            normalize_service_time_input("08:30 01.01.2025"),
        )

    def test_service_period_fields_insert_separators_automatically(self):
        window = MainWindow()
        try:
            window.ed_service_from.insert("0830")
            window.ed_service_to.insert("1045")
            self.assertEqual("08:30", window.ed_service_from.text())
            self.assertEqual("10:45", window.ed_service_to.text())
            self.assertEqual(
                "00:00;_",
                window.ed_service_from.inputMask(),
            )
            self.assertEqual(
                "__:__",
                window.ed_service_from.placeholderText(),
            )
        finally:
            window.close()

    def test_service_period_keeps_manual_entry_without_visible_arrows(self):
        window = MainWindow()
        try:
            window.ed_service_from.insert("0915")
            self.assertEqual("09:15", window.ed_service_from.text())
            self.assertTrue(window.ed_service_from.picker_button.isHidden())
            self.assertTrue(window.ed_service_to.picker_button.isHidden())
        finally:
            window.close()

    def test_surname_only_button_creates_both_blank_template_types(self):
        self.assertEqual(
            "Лазутин Д. М.",
            executor_name_with_initials("123! Лазутин-42 Д. М."),
        )
        self.assertEqual(
            "ПетровСидоров А.",
            executor_name_with_initials("Петров-Сидоров А."),
        )
        self.assertEqual("Лазутин", executor_name_with_initials("Лазутин"))

        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "Acts"
            history = Path(directory) / "History"
            window = MainWindow()
            try:
                for listed_executor in window.data_lists["executors"]:
                    with self.subTest(listed_executor=listed_executor):
                        self.assertEqual(
                            listed_executor,
                            executor_name_with_initials(listed_executor),
                        )
                self.assertIsNotNone(window.btn_surname_act.menu())
                self.assertEqual(
                    ["АБП / МКТФ", "Валидатор / МИД"],
                    [action.text() for action in window.btn_surname_act.menu().actions()],
                )
                window.app_settings["preview_before_save"] = False
                window.cb_executor.setEditText("Лазутин Д. М. 123!")
                window.ed_date.setText("07.10.2026")
                window.cb_location.setEditText("ОСТАНКИНО")
                window.cb_model.setEditText("МКТФ")
                window.ed_serial.setText("52840")
                window.cb_work.setEditText("НЕ ДОЛЖНО ПОПАСТЬ В АКТ")

                with (
                    patch("app.OUTPUT_DIR", str(output)),
                    patch("app.HISTORY_DIR", str(history)),
                    patch("app.append_act_to_workbook") as append,
                    patch.object(QMessageBox, "exec", return_value=0),
                    patch.object(QMessageBox, "clickedButton", return_value=None),
                ):
                    window.create_surname_only_act(TEMPLATE_TYPE1)
                    window.create_surname_only_act(TEMPLATE_TYPE2)

                append.assert_not_called()
                documents = sorted(output.glob("*.docx"))
                self.assertEqual(2, len(documents))
                self.assertEqual(2, len(list(history.glob("*.json"))))
                for path in documents:
                    document = Document(path)
                    text = "\n".join(
                        [paragraph.text for paragraph in document.paragraphs]
                        + [
                            paragraph.text
                            for table in document.tables
                            for cell in table._cells
                            for paragraph in cell.paragraphs
                        ]
                    )
                    self.assertIn("Лазутин Д. М.", text)
                    self.assertNotIn("123", text)
                    self.assertNotIn("07.10.2026", text)
                    self.assertNotIn("ОСТАНКИНО", text)
                    self.assertNotIn("52840", text)
                    self.assertNotIn("НЕ ДОЛЖНО ПОПАСТЬ В АКТ", text)
                    self.assertNotIn("{EXECUTOR}", text)
                    equipment_value_cells = document.tables[0].rows[9].cells[1:]
                    self.assertTrue(equipment_value_cells)
                    self.assertTrue(
                        all(not cell.text.strip() for cell in equipment_value_cells)
                    )
            finally:
                window.close()

    def test_empty_masked_service_period_is_optional(self):
        window = MainWindow()
        try:
            self.assertEqual(
                ("_______часов «__»____202__года",) * 2,
                validate_service_period(
                    window.ed_service_from.text(),
                    window.ed_service_to.text(),
                    "17.08.2026",
                ),
            )
            data = window.current_act_data()
            self.assertEqual("", data["service_from"])
            self.assertEqual("", data["service_to"])
            window.ed_service_from.insert("0830")
            with self.assertRaises(ValueError):
                validate_service_period(
                    window.ed_service_from.text(),
                    window.ed_service_to.text(),
                    "17.08.2026",
                )
        finally:
            window.close()

    def test_packaged_auto_update_check_is_enabled(self):
        config_path = Path(__file__).resolve().parents[1] / "update_config.json"
        with config_path.open("r", encoding="utf-8") as stream:
            config = json.load(stream)
        self.assertIs(config.get("check_on_start"), True)

    def test_release_version_is_1_0_22(self):
        self.assertEqual("1.0.22", __version__)

    def test_date_field_keeps_manual_entry_and_uses_visible_picker_button(self):
        window = MainWindow()
        try:
            window.ed_date.setText("17.08.2026")
            self.assertEqual("17.08.2026", window.ed_date.text())
            self.assertEqual(30, window.ed_date.calendar_button.width())
            self.assertFalse(window.ed_date.calendar_button.icon().isNull())
            window.ed_date.show_calendar()
            menu = window.ed_date._calendar_menu
            self.assertIsNotNone(menu)
            window.ed_date._select_calendar_date(QDate(2026, 8, 19), menu)
            self.assertEqual("19.08.2026", window.ed_date.text())
        finally:
            window.close()

    def test_factory_number_fills_equipment_and_station_from_registry(self):
        window = MainWindow()
        try:
            self.assertEqual(382, len(window.equipment_registry))
            self.assertEqual("АЛАБУШЕВО", window.data_lists["locations"][-1])
            self.assertNotIn("Зеленоград-Крюково", window.data_lists["locations"])
            self.assertTrue(
                all(
                    record["station"] != "Зеленоград-Крюково"
                    for record in window.equipment_registry.values()
                )
            )
            self.assertTrue(
                any(
                    record["station"] == "КРЮКОВО"
                    for record in window.equipment_registry.values()
                )
            )
            self.assertEqual(0, window.ed_serial.count())

            window.cb_location.setEditText("ОСТАНКИНО")
            self.assertEqual(14, window.ed_serial.count())

            window.cb_model.setEditText("МКТФ")
            self.assertEqual(3, window.ed_serial.count())
            mktf_index = window.ed_serial.findText("52840")
            self.assertGreaterEqual(mktf_index, 0)
            self.assertEqual(-1, window.ed_serial.findText("180300088"))
            self.assertEqual(-1, window.ed_serial.findText("53886"))

            window.ed_serial.setCurrentIndex(mktf_index)
            self.assertEqual("МКТФ", window.cb_model.text())
            self.assertEqual("ОСТАНКИНО", window.cb_location.text())

            window.cb_model.setEditText("АБП-БН")
            self.assertEqual(9, window.ed_serial.count())
            abp_bn_index = window.ed_serial.findText("180300088")
            self.assertGreaterEqual(abp_bn_index, 0)
            self.assertEqual(-1, window.ed_serial.findText("52840"))
            window.ed_serial.setEditText("")
            window.ed_serial.textActivated.emit("180300088")
            self.assertEqual("АБП-БН", window.cb_model.text())
            self.assertEqual("ОСТАНКИНО", window.cb_location.text())

            window.cb_model.setEditText("АБП-09")
            self.assertEqual(2, window.ed_serial.count())
            self.assertGreaterEqual(window.ed_serial.findText("53886"), 0)
            self.assertEqual(-1, window.ed_serial.findText("180300088"))

            window.cb_model.setEditText("")
            window.cb_location.setEditText("НОВОПОДРЕЗКОВО")
            self.assertEqual(12, window.ed_serial.count())
            self.assertGreaterEqual(window.ed_serial.findText("52881"), 0)
            self.assertEqual(-1, window.ed_serial.findText("52817"))
            window.cb_location.setEditText("Другая станция")
            window.cb_model.setEditText("МИД")
            window.ed_serial.setText("52881")
            self.assertTrue(window.apply_equipment_for_serial())
            self.assertEqual("НОВОПОДРЕЗКОВО", window.cb_location.text())
            self.assertEqual("МКТФ", window.cb_model.text())
            self.assertIn("52881", window.lbl_status.text())

            window.ed_serial.setText("999999")
            self.assertFalse(window.apply_equipment_for_serial())
            self.assertEqual("НОВОПОДРЕЗКОВО", window.cb_location.text())
            self.assertEqual("МКТФ", window.cb_model.text())
        finally:
            window.close()

    def test_source_equipment_names_use_application_model_names(self):
        expected_aliases = {
            "АБП-БН": "АБП-БН",
            "АПБ": "АБП-09",
            "АН-15": "АН-15",
            "БПА-20-БН1": "БПА-20-БН1",
            "АПБ-09-М3": "АБП-09-М3",
            "МКТФ": "МКТФ",
        }
        for source_name, application_name in expected_aliases.items():
            with self.subTest(source_name=source_name):
                self.assertEqual(
                    application_name,
                    canonical_equipment_model(source_name),
                )

        window = MainWindow()
        try:
            examples = {
                "180300002": "АБП-БН",
                "52366": "АБП-09",
                "56745": "АН-15",
                "121": "АБП-09-М3",
                "793": "БПА-20-БН1",
                "52795": "МКТФ",
            }
            for serial, expected_model in examples.items():
                with self.subTest(serial=serial):
                    self.assertTrue(window.apply_equipment_for_serial(serial))
                    self.assertEqual(expected_model, window.cb_model.text())
        finally:
            window.close()

    def test_fullscreen_layout_keeps_rows_and_action_bar_compact(self):
        window = MainWindow()
        try:
            window.resize(2560, 1440)
            window.show()
            self.application.processEvents()
            self.assertEqual(42, window.ed_date.height())
            self.assertLessEqual(window.action_bar.height(), 70)
            self.assertLessEqual(window.cb_executor.height(), 44)
            self.assertEqual(42, window.sp_qty.height())
            self.assertEqual(30, window.sp_qty.step_column.width())
            self.assertFalse(window.sp_qty.up_button.icon().isNull())
            self.assertFalse(window.sp_qty.down_button.icon().isNull())
            self.assertEqual(QSize(12, 12), window.sp_qty.up_button.iconSize())
            self.assertEqual(QSize(12, 12), window.sp_qty.down_button.iconSize())
            self.assertEqual(
                window.sp_qty.step_column.rect().center().x(),
                window.sp_qty.up_button.geometry().center().x(),
            )
            self.assertEqual(
                window.sp_qty.step_column.rect().center().x(),
                window.sp_qty.down_button.geometry().center().x(),
            )
            self.assertLessEqual(
                window.sp_qty.step_column.geometry().right(),
                window.sp_qty.container.rect().right(),
            )
        finally:
            window.close()

    def test_excluded_executor_is_removed_and_cannot_be_used(self):
        self.assertTrue(is_excluded_executor("Чепель А. А."))
        self.assertTrue(is_excluded_executor("чепель аа"))
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "executors.txt"
            path.write_text("Петров К. Н.\nЧепель А. А.\n", encoding="utf-8")
            self.assertEqual(["Петров К. Н."], read_lines(str(path)))
            write_lines(str(path), ["Чепель А.А.", "Лазутин Д. М."])
            self.assertEqual(["Лазутин Д. М."], read_lines(str(path)))

        window = MainWindow()
        try:
            self.assertFalse(any(
                is_excluded_executor(window.cb_executor.itemText(index))
                for index in range(window.cb_executor.count())
            ))
            window.cb_executor.setEditText("Чепель А. А.")
            self.assertEqual("", window.current_act_data()["executor"])
            self.assertEqual("", window.cb_executor.text())
        finally:
            window.close()

    def test_preview_menu_setting_is_loaded_and_persisted(self):
        with (
            patch("app.load_app_settings", return_value={"preview_before_save": False}),
            patch("app.save_app_settings") as save_settings,
        ):
            window = MainWindow()
            try:
                self.assertFalse(window.act_preview_enabled.isChecked())
                window.act_preview_enabled.setChecked(True)
                save_settings.assert_called_once_with({"preview_before_save": True})
            finally:
                window.close()

    def test_disabled_preview_saves_without_opening_dialog(self):
        class UnexpectedPreview:
            def __init__(self, *_args, **_kwargs):
                raise AssertionError("Preview dialog must not be created")

        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "Acts"
            window = MainWindow()
            try:
                window.app_settings["preview_before_save"] = False
                window.cb_location.setEditText("Химки")
                window.cb_model.setEditText("МКТФ")
                window.cb_template.setCurrentIndex(1)

                excel_result = SimpleNamespace(sheet_name="Август 26", row_number=2)
                with (
                    patch("app.OUTPUT_DIR", str(output)),
                    patch("app.HISTORY_DIR", str(Path(directory) / "History")),
                    patch("app.ActPreviewDialog", UnexpectedPreview),
                    patch.object(window, "active_excel_workbook", return_value="Tables.xlsx"),
                    patch("app.append_act_to_workbook", return_value=excel_result),
                    patch.object(QMessageBox, "exec", return_value=0),
                    patch.object(QMessageBox, "clickedButton", return_value=None),
                ):
                    window.create_act()

                self.assertEqual(1, len(list(output.glob("*.docx"))))
            finally:
                window.close()

    def test_visual_preview_loads_a_real_pdf_page(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            pdf_path = root / "preview.pdf"
            dummy_docx = root / "preview.docx"
            dummy_docx.write_bytes(b"test")

            writer = QPdfWriter(str(pdf_path))
            painter = QPainter(writer)
            painter.drawText(QRectF(100, 100, 800, 200), "Act preview smoke test")
            painter.end()
            del painter
            del writer

            release_renderer = threading.Event()

            def delayed_renderer(_path):
                release_renderer.wait(2)
                return str(pdf_path)

            with patch("app.render_docx_to_pdf", side_effect=delayed_renderer):
                dialog = ActPreviewDialog(str(dummy_docx))
                try:
                    self.assertTrue(dialog.render_thread.isRunning())
                    self.assertIsNone(dialog.pdf_document)
                    self.assertFalse(dialog.open_button.isEnabled())
                    release_renderer.set()
                    self.assertTrue(dialog.render_thread.wait(3000))
                    self.application.processEvents()
                    self.assertEqual(1, dialog.pdf_document.pageCount())
                    self.assertIs(dialog.pdf_document, dialog.preview.document())
                    self.assertTrue(dialog.open_button.isEnabled())
                    self.assertTrue(dialog.save_button.isEnabled())
                finally:
                    dialog.reject()

    def test_form_creates_docx_and_passes_derived_excel_values(self):
        class AcceptedPreview:
            def __init__(self, _path, _parent=None):
                pass

            def exec(self):
                return QDialog.Accepted

        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "Acts"
            window = MainWindow()
            try:
                window.ed_date.setText("08.08.2026")
                window.ed_service_from.setText("08:30")
                window.ed_service_to.setText("10:45")
                window.cb_executor.setEditText("Петров")
                window.cb_location.setEditText("Химки")
                window.cb_model.setEditText("МИД")
                window.ed_serial.setText("999999")
                window.cb_work.setEditText("Замена блока питания")
                window.cb_issues.setEditText("Сбой")
                window.cb_done_work.setEditText("Блок заменён")
                window.cb_materials.setEditText("Комплект МТТПК")
                window.sp_qty.setValue(1)
                window.cb_template.setCurrentIndex(0)
                self.assertEqual(
                    "Validator_MID.docx",
                    Path(window.selected_template(window.cb_model.text())).name,
                )

                excel_result = SimpleNamespace(sheet_name="Август 26", row_number=2)
                with (
                    patch("app.OUTPUT_DIR", str(output)),
                    patch("app.HISTORY_DIR", str(Path(directory) / "History")),
                    patch("app.ActPreviewDialog", AcceptedPreview),
                    patch.object(window, "active_excel_workbook", return_value="Tables.xlsx"),
                    patch("app.append_act_to_workbook", return_value=excel_result) as append,
                    patch.object(QMessageBox, "exec", return_value=0),
                    patch.object(QMessageBox, "clickedButton", return_value=None),
                ):
                    window.create_act()

                documents = list(output.glob("*.docx"))
                self.assertEqual(1, len(documents))
                self.assertEqual(1, len(list((Path(directory) / "History").glob("*.json"))))
                values = append.call_args.args[1]
                self.assertEqual("Крюковский", values["area"])
                self.assertEqual("выданно МТППК", values["notes"])
                self.assertEqual("Химки", values["location"])
                generated = Document(documents[0])
                all_text = "\n".join(
                    paragraph.text
                    for table in generated.tables
                    for cell in table._cells
                    for paragraph in cell.paragraphs
                )
                self.assertIn("08:30 часов «08» августа 2026 года", all_text)
                self.assertIn("10:45 часов «08» августа 2026 года", all_text)
            finally:
                window.close()

    def test_batch_panel_creates_all_queued_acts_without_preview(self):
        class UnexpectedPreview:
            def __init__(self, *_args, **_kwargs):
                raise AssertionError("Batch creation must not open previews")

        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "Acts"
            window = MainWindow()
            try:
                window.act_batch_panel.setChecked(True)
                self.assertFalse(window.batch_panel.isHidden())
                window.cb_location.setEditText("Химки")
                window.cb_model.setEditText("МКТФ")
                window.ed_serial.setText("1001")
                window.cb_template.setCurrentIndex(1)
                window.add_current_act_to_batch()
                window.ed_serial.setText("1002")
                window.add_current_act_to_batch()
                self.assertEqual(2, window.batch_table.rowCount())

                excel_result = SimpleNamespace(sheet_name="Август 26", row_number=2)
                with (
                    patch("app.OUTPUT_DIR", str(output)),
                    patch("app.HISTORY_DIR", str(Path(directory) / "History")),
                    patch("app.ActPreviewDialog", UnexpectedPreview),
                    patch.object(window, "active_excel_workbook", return_value="Tables.xlsx"),
                    patch("app.append_act_to_workbook", return_value=excel_result) as append,
                    patch.object(QMessageBox, "information", return_value=QMessageBox.Ok),
                ):
                    window.create_batch_acts()

                self.assertEqual(2, len(list(output.glob("*.docx"))))
                self.assertEqual(2, append.call_count)
                self.assertEqual(0, window.batch_table.rowCount())
            finally:
                window.close()

    def test_batch_row_can_be_edited_and_duplicated(self):
        window = MainWindow()
        try:
            window.act_batch_panel.setChecked(True)
            window.cb_location.setEditText("Химки")
            window.cb_model.setEditText("МКТФ")
            window.ed_serial.setText("1001")
            window.cb_template.setCurrentIndex(1)
            window.add_current_act_to_batch()

            window.edit_batch_act(0, 3)
            self.assertEqual("1001", window.ed_serial.text())
            self.assertEqual("Сохранить изменения", window.btn_batch_add.text())
            window.ed_serial.setText("2002")
            window.add_current_act_to_batch()
            self.assertEqual("2002", window.batch_row_data(0)["serial"])

            window.batch_table.selectRow(0)
            window.duplicate_selected_batch_act()
            self.assertEqual(2, window.batch_table.rowCount())
            self.assertEqual("В очереди: 2", window.batch_count_label.text())
        finally:
            window.close()

    def test_batch_queue_clear_requires_confirmation(self):
        window = MainWindow()
        try:
            window.act_batch_panel.setChecked(True)
            window.cb_location.setEditText("Химки")
            window.cb_model.setEditText("МКТФ")
            window.ed_serial.setText("1001")
            window.cb_template.setCurrentIndex(1)
            window.add_current_act_to_batch()
            with patch.object(QMessageBox, "question", return_value=QMessageBox.No):
                window.clear_batch_acts()
            self.assertEqual(1, window.batch_table.rowCount())
            with patch.object(QMessageBox, "question", return_value=QMessageBox.Yes):
                window.clear_batch_acts()
            self.assertEqual(0, window.batch_table.rowCount())
        finally:
            window.close()

    def test_history_opens_real_docx_and_can_load_form(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            acts = root / "Acts"
            history = root / "History"
            acts.mkdir()
            document = acts / "history.docx"
            document.write_bytes(b"docx")
            data = {
                "date": "11.08.2026", "executor": "Петров",
                "location": "Химки", "model": "МКТФ", "serial": "7788",
                "work": "Проверка", "issues": "", "done_work": "Готово",
                "materials": "", "qty": 0, "template": "",
            }
            from history_store import save_history_record
            save_history_record(history, root, document, data)
            window = MainWindow()
            try:
                with (
                    patch("app.HISTORY_DIR", str(history)),
                    patch("app.APP_DIR", str(root)),
                    patch("app.OUTPUT_DIR", str(acts)),
                ):
                    window.act_history_panel.setChecked(True)
                    self.assertFalse(window.act_batch_panel.isChecked())
                    self.assertEqual(1, window.history_table.rowCount())
                    window.history_table.selectRow(0)
                    window.load_selected_history()
                    self.assertEqual("7788", window.ed_serial.text())
                    with patch("app.open_file") as open_document:
                        window.open_selected_history()
                        open_document.assert_called_once_with(str(document.resolve()))
            finally:
                window.close()

    def test_history_clear_keeps_docx(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            acts = root / "Acts"
            history = root / "History"
            acts.mkdir()
            document = acts / "history.docx"
            document.write_bytes(b"docx")
            from history_store import save_history_record
            save_history_record(
                history, root, document,
                {"date": "11.08.2026", "serial": "7788"},
            )
            window = MainWindow()
            try:
                with (
                    patch("app.HISTORY_DIR", str(history)),
                    patch("app.APP_DIR", str(root)),
                    patch("app.OUTPUT_DIR", str(acts)),
                    patch.object(QMessageBox, "question", return_value=QMessageBox.Yes),
                ):
                    window.act_history_panel.setChecked(True)
                    self.assertEqual(1, window.history_table.rowCount())
                    window.clear_history()
                    self.assertEqual(0, window.history_table.rowCount())
                    self.assertTrue(document.is_file())
            finally:
                window.close()

    def test_delete_selected_history_moves_only_selected_docx(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            acts = root / "Acts"
            history = root / "History"
            acts.mkdir()
            selected = acts / "selected.docx"
            untouched = acts / "untouched.docx"
            selected.write_bytes(b"selected")
            untouched.write_bytes(b"untouched")
            from history_store import save_history_record
            save_history_record(history, root, selected, {"serial": "1"})
            save_history_record(history, root, untouched, {"serial": "2"})
            window = MainWindow()
            try:
                with (
                    patch("app.HISTORY_DIR", str(history)),
                    patch("app.APP_DIR", str(root)),
                    patch("app.OUTPUT_DIR", str(acts)),
                    patch.object(QMessageBox, "question", return_value=QMessageBox.Yes),
                    patch("app.move_file_to_recycle_bin", side_effect=lambda path: Path(path).unlink()),
                ):
                    window.act_history_panel.setChecked(True)
                    for row in range(window.history_table.rowCount()):
                        if window.history_table.item(row, 5).text() == selected.name:
                            window.history_table.selectRow(row)
                            break
                    window.delete_selected_history()
                    self.assertFalse(selected.exists())
                    self.assertTrue(untouched.exists())
                    self.assertEqual(1, window.history_table.rowCount())
                    self.assertEqual(untouched.name, window.history_table.item(0, 5).text())
            finally:
                window.close()

    def test_delete_selected_history_can_be_cancelled(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            acts = root / "Acts"
            history = root / "History"
            acts.mkdir()
            document = acts / "kept.docx"
            document.write_bytes(b"docx")
            from history_store import save_history_record
            save_history_record(history, root, document, {"serial": "1"})
            window = MainWindow()
            try:
                with (
                    patch("app.HISTORY_DIR", str(history)),
                    patch("app.APP_DIR", str(root)),
                    patch("app.OUTPUT_DIR", str(acts)),
                    patch.object(QMessageBox, "question", return_value=QMessageBox.No),
                    patch("app.move_file_to_recycle_bin") as recycle,
                ):
                    window.act_history_panel.setChecked(True)
                    window.history_table.selectRow(0)
                    window.delete_selected_history()
                    self.assertTrue(document.exists())
                    recycle.assert_not_called()
            finally:
                window.close()


if __name__ == "__main__":
    unittest.main()
