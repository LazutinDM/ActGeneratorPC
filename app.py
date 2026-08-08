"""ActGeneratorPC user interface and document-generation workflow.

Folder paths live in app_paths.py; update/rollback code lives in
update_manager.py.  Keep this file focused on UI actions and DOCX generation.
"""

import os
import re
import sys
import json
import datetime
import subprocess
import ctypes
import shutil
import tempfile
from typing import List, Dict
from functools import partial

from update_manager import UpdateManager, confirm_healthy_startup
from PySide6.QtCore import (
    Qt, QUrl, QSignalBlocker, QStringListModel, QEvent, QTimer,
    QByteArray, QBuffer, QIODevice, QThread,
)
from PySide6.QtPdf import QPdfDocument
from PySide6.QtPdfWidgets import QPdfView
from PySide6.QtGui import (
    QAction, QPalette, QColor, QDesktopServices, QPixmap, QPainter, QPen, QIcon,
    QPainterPath, QRegion, QFont
)
from PySide6.QtWidgets import (
    QApplication, QMainWindow, QWidget, QVBoxLayout, QHBoxLayout, QLabel, QLineEdit,
    QPushButton, QMessageBox, QDialog, QListWidget, QListWidgetItem,
    QAbstractItemView, QInputDialog, QSpinBox, QFormLayout,
    QStyle, QComboBox, QCompleter, QToolButton, QMenu, QFrame,
    QGraphicsDropShadowEffect, QFileDialog
)

from docx import Document
from docx_preview import PreviewRenderError, render_docx_to_pdf
from excel_export import (
    ExcelExportError,
    append_act_to_workbook,
    area_for_station,
    ensure_embedded_excel_workbook,
    load_excel_path,
    note_for_materials,
    resolve_excel_workbook,
    save_excel_path,
)

from app_paths import (
    APP_DIR, APP_SETTINGS_CONFIG, CONFIG_DIR, DATA_DIR, DOCUMENTATION_DIR,
    EMBEDDED_EXCEL_WORKBOOK, EXCEL_EXPORT_CONFIG, EXCEL_TEMPLATE,
    EXCEL_WORK_DIR, FILES, ICONS_DIR,
    ICONS_GEN_DIR, ICON_APP, ICON_ARROW_DOWN, ICON_CALENDAR, ICON_DOCX,
    ICON_DONE_WORK, ICON_EQUIPMENT, ICON_EXCEL, ICON_ISSUES, ICON_LISTS,
    ICON_MATERIALS, ICON_QTY, ICON_SERIAL, ICON_SERVICES, ICON_STATION,
    ICON_USER, ICON_XLS, LEGACY_OUTPUT_DIR, OTHER_DIR, OUTPUT_DIR,
    SPIN_DOWN_DARK_PNG, SPIN_DOWN_LIGHT_PNG, SPIN_UP_DARK_PNG,
    SPIN_UP_LIGHT_PNG, TEMPLATE_TYPE1, TEMPLATE_TYPE2, TEMPLATES_DIR,
    VARIABLES_DIR,
)


INVALID_FILENAME_RE = re.compile(r'[\\/:*?"<>|]+')
URL_SCHEME_RE = re.compile(r"^[a-zA-Z][a-zA-Z0-9+\-.]*://")
SEARCH_CLEAN_RE = re.compile(r"[^0-9a-zа-я]+")


# -----------------------------
# Helpers
# -----------------------------
def ensure_dirs():
    os.makedirs(DATA_DIR, exist_ok=True)
    os.makedirs(TEMPLATES_DIR, exist_ok=True)
    os.makedirs(VARIABLES_DIR, exist_ok=True)
    os.makedirs(OTHER_DIR, exist_ok=True)
    os.makedirs(CONFIG_DIR, exist_ok=True)
    os.makedirs(DOCUMENTATION_DIR, exist_ok=True)
    os.makedirs(EXCEL_WORK_DIR, exist_ok=True)
    os.makedirs(OUTPUT_DIR, exist_ok=True)
    os.makedirs(ICONS_GEN_DIR, exist_ok=True)
    os.makedirs(ICONS_DIR, exist_ok=True)
    migrate_legacy_portable_layout()


def migrate_legacy_portable_layout():
    """Move user documents and remove obsolete root items after an update."""
    # Older versions stored variable lists directly in Data.
    for name in os.listdir(DATA_DIR):
        source = os.path.join(DATA_DIR, name)
        if not os.path.isfile(source) or not name.lower().endswith(".txt"):
            continue
        destination = os.path.join(VARIABLES_DIR, name)
        try:
            if os.path.exists(destination):
                os.remove(source)
            else:
                os.replace(source, destination)
        except OSError:
            if not os.path.exists(destination):
                shutil.copy2(source, destination)
                os.remove(source)

    legacy_tables = os.path.join(DATA_DIR, "tables.json")
    current_tables = os.path.join(CONFIG_DIR, "tables.json")
    if os.path.isfile(legacy_tables):
        try:
            if os.path.exists(current_tables):
                os.remove(legacy_tables)
            else:
                os.replace(legacy_tables, current_tables)
        except OSError:
            if not os.path.exists(current_tables):
                shutil.copy2(legacy_tables, current_tables)
                os.remove(legacy_tables)

    legacy_generated_icons = os.path.join(DATA_DIR, "_icons")
    if os.path.isdir(legacy_generated_icons):
        for name in os.listdir(legacy_generated_icons):
            source = os.path.join(legacy_generated_icons, name)
            destination = os.path.join(ICONS_GEN_DIR, name)
            if os.path.isfile(source) and not os.path.exists(destination):
                try:
                    os.replace(source, destination)
                except OSError:
                    shutil.copy2(source, destination)
                    os.remove(source)
        try:
            os.rmdir(legacy_generated_icons)
        except OSError:
            pass

    legacy_output_dirs = (
        LEGACY_OUTPUT_DIR,
        os.path.join(APP_DIR, "\u0410\u043a\u0442\u044b"),
    )
    for legacy_output_dir in legacy_output_dirs:
        if not os.path.isdir(legacy_output_dir):
            continue
        for name in os.listdir(legacy_output_dir):
            source = os.path.join(legacy_output_dir, name)
            destination = os.path.join(OUTPUT_DIR, name)
            if not os.path.exists(destination):
                try:
                    os.replace(source, destination)
                except OSError:
                    if os.path.isdir(source):
                        shutil.copytree(source, destination)
                        shutil.rmtree(source)
                    else:
                        shutil.copy2(source, destination)
                        os.remove(source)
        try:
            os.rmdir(legacy_output_dir)
        except OSError:
            pass

    if not getattr(sys, "frozen", False):
        return

    # Normalize the spelling of portable user-data folders. PyInstaller runtime
    # files also live in Data, so unknown entries must never be deleted here.
    expected_data_folders = {
        name.casefold(): name
        for name in ("Icons", "Other", "Templates", "Variables")
    }
    # Windows paths are case-insensitive, but Explorer keeps the spelling used
    # by an old release. Rename through a temporary name to normalize it.
    for actual_name in os.listdir(DATA_DIR):
        canonical_name = expected_data_folders.get(actual_name.casefold())
        if not canonical_name or actual_name == canonical_name:
            continue
        actual_path = os.path.join(DATA_DIR, actual_name)
        temporary_path = os.path.join(DATA_DIR, f".{canonical_name}.rename")
        canonical_path = os.path.join(DATA_DIR, canonical_name)
        try:
            os.replace(actual_path, temporary_path)
            os.replace(temporary_path, canonical_path)
        except OSError:
            pass

    for legacy_folder in ("_internal", "Templates", "icons", "_icons", "variables"):
        legacy_path = os.path.join(APP_DIR, legacy_folder)
        if os.path.isdir(legacy_path):
            shutil.rmtree(legacy_path, ignore_errors=True)

    legacy_config = os.path.join(APP_DIR, "update_config.json")
    new_config = os.path.join(CONFIG_DIR, "update_config.json")
    if os.path.isfile(legacy_config) and os.path.isfile(new_config):
        try:
            os.remove(legacy_config)
        except OSError:
            pass


def _make_placeholder_icon_png(path: str, label: str, bg: QColor):
    if os.path.exists(path):
        return
    pm = QPixmap(64, 64)
    pm.fill(Qt.transparent)

    p = QPainter(pm)
    p.setRenderHint(QPainter.Antialiasing, True)

    p.setPen(Qt.NoPen)
    p.setBrush(bg)
    p.drawRoundedRect(4, 4, 56, 56, 14, 14)

    p.setPen(QColor("#ffffff"))
    f = QFont()
    f.setBold(True)
    f.setPointSize(16)
    p.setFont(f)
    p.drawText(pm.rect(), Qt.AlignCenter, (label or "?")[:3].upper())

    p.end()
    pm.save(path, "PNG")


def ensure_default_icons():
    _make_placeholder_icon_png(ICON_ARROW_DOWN, "V", QColor("#4B5563"))
    _make_placeholder_icon_png(ICON_CALENDAR, "CAL", QColor("#10B981"))
    _make_placeholder_icon_png(ICON_DONE_WORK, "OK", QColor("#22C55E"))
    _make_placeholder_icon_png(ICON_EXCEL, "XL", QColor("#16A34A"))
    _make_placeholder_icon_png(ICON_ISSUES, "!", QColor("#F97316"))
    _make_placeholder_icon_png(ICON_MATERIALS, "MAT", QColor("#0EA5E9"))
    _make_placeholder_icon_png(ICON_QTY, "QTY", QColor("#6366F1"))
    _make_placeholder_icon_png(ICON_SERIAL, "SN", QColor("#8B5CF6"))
    _make_placeholder_icon_png(ICON_SERVICES, "SRV", QColor("#06B6D4"))
    _make_placeholder_icon_png(ICON_STATION, "ST", QColor("#64748B"))
    _make_placeholder_icon_png(ICON_USER, "USR", QColor("#3B82F6"))
    _make_placeholder_icon_png(ICON_XLS, "XLS", QColor("#16A34A"))
    _make_placeholder_icon_png(ICON_EQUIPMENT, "EQ", QColor("#A855F7"))
    _make_placeholder_icon_png(ICON_LISTS, "LST", QColor("#14B8A6"))
    _make_placeholder_icon_png(ICON_DOCX, "DOC", QColor("#2563EB"))


def make_spin_arrow_png(path: str, color_hex: str, direction: str):
    if os.path.exists(path):
        return
    pm = QPixmap(18, 18)
    pm.fill(Qt.transparent)

    p = QPainter(pm)
    p.setRenderHint(QPainter.Antialiasing, True)
    pen = QPen(QColor(color_hex))
    pen.setWidthF(3.2)
    pen.setCapStyle(Qt.RoundCap)
    pen.setJoinStyle(Qt.RoundJoin)
    p.setPen(pen)

    if direction == "up":
        p.drawLine(4, 11, 9, 6)
        p.drawLine(9, 6, 14, 11)
    else:
        p.drawLine(4, 7, 9, 12)
        p.drawLine(9, 12, 14, 7)

    p.end()
    pm.save(path, "PNG")


def ensure_arrow_icons():
    make_spin_arrow_png(SPIN_UP_LIGHT_PNG, "#111111", "up")
    make_spin_arrow_png(SPIN_DOWN_LIGHT_PNG, "#111111", "down")
    make_spin_arrow_png(SPIN_UP_DARK_PNG, "#FFFFFF", "up")
    make_spin_arrow_png(SPIN_DOWN_DARK_PNG, "#FFFFFF", "down")


def resolve_portable_path(path: str) -> str:
    path = (path or "").strip()
    if not path or os.path.isabs(path):
        return path
    return os.path.normpath(os.path.join(APP_DIR, path))


def portable_stored_path(path: str) -> str:
    path = (path or "").strip()
    if not path:
        return ""
    try:
        relative = os.path.relpath(os.path.abspath(path), APP_DIR)
    except ValueError:
        return path
    if relative == ".." or relative.startswith(".." + os.sep):
        return path
    return relative.replace("\\", "/")


def icon_or_empty(path: str) -> QIcon:
    resolved = resolve_portable_path(path)
    return QIcon(resolved) if resolved and os.path.exists(resolved) else QIcon()


def pixmap_or_empty(path: str, size: int = 18) -> QPixmap:
    path = resolve_portable_path(path)
    if not path or not os.path.exists(path):
        return QPixmap()
    pm = QPixmap(path)
    if pm.isNull():
        return QPixmap()
    return pm.scaled(size, size, Qt.KeepAspectRatio, Qt.SmoothTransformation)


def make_form_label(text: str, icon_path: str) -> QWidget:
    w = QWidget()
    lay = QHBoxLayout(w)
    lay.setContentsMargins(0, 0, 0, 0)
    lay.setSpacing(8)

    ico = QLabel()
    pm = pixmap_or_empty(icon_path, 18)
    if not pm.isNull():
        ico.setPixmap(pm)
    ico.setFixedWidth(20)

    lbl = QLabel(text)
    lbl.setAlignment(Qt.AlignVCenter | Qt.AlignLeft)

    lay.addWidget(ico)
    lay.addWidget(lbl)
    lay.addStretch(1)
    return w


def read_lines(path: str) -> List[str]:
    if not os.path.exists(path):
        return []
    with open(path, "r", encoding="utf-8") as f:
        return [text for line in f if (text := line.strip())]


def write_lines(path: str, items: List[str]):
    os.makedirs(os.path.dirname(path), exist_ok=True)
    clean = [x.strip() for x in items if x and x.strip()]
    with open(path, "w", encoding="utf-8") as f:
        f.write("\n".join(clean) + ("\n" if clean else ""))


def safe_filename(s: str) -> str:
    s = (s or "").strip()
    return INVALID_FILENAME_RE.sub("_", s)[:180]


def open_folder(path: str):
    if os.name == "nt":
        os.startfile(path)
    elif sys.platform == "darwin":
        subprocess.run(["open", path])
    else:
        subprocess.run(["xdg-open", path])


def normalize_url_or_path(s: str) -> QUrl:
    s = (s or "").strip()
    if not s:
        return QUrl()
    if URL_SCHEME_RE.match(s):
        return QUrl(s)
    if os.path.exists(s):
        return QUrl.fromLocalFile(os.path.abspath(s))
    if s.startswith("www."):
        return QUrl("https://" + s)
    return QUrl(s)


def open_any(link_or_path: str):
    url = normalize_url_or_path(link_or_path)
    if not url.isValid() or url.isEmpty():
        QMessageBox.warning(None, "Открытие", "Пустая/неверная ссылка или путь.")
        return
    QDesktopServices.openUrl(url)


def open_file(path: str):
    if not path:
        return
    QDesktopServices.openUrl(QUrl.fromLocalFile(os.path.abspath(path)))


def normalize_search(text: str) -> str:
    return SEARCH_CLEAN_RE.sub("", (text or "").strip().lower().replace("ё", "е"))


def replace_in_paragraph(paragraph, mapping: Dict[str, str]):
    full = "".join(run.text for run in paragraph.runs)
    if not full:
        return
    changed = False
    for key, val in mapping.items():
        if key in full:
            full = full.replace(key, val)
            changed = True
    if changed:
        if paragraph.runs:
            paragraph.runs[0].text = full
            for r in paragraph.runs[1:]:
                r.text = ""
        else:
            paragraph.add_run(full)


def replace_placeholders_docx(doc: Document, mapping: Dict[str, str]):
    for p in doc.paragraphs:
        replace_in_paragraph(p, mapping)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    replace_in_paragraph(p, mapping)


# -----------------------------
# Links storage (JSON)
# -----------------------------
def load_links() -> List[Dict[str, str]]:
    if not os.path.exists(FILES["links"]):
        return []
    try:
        with open(FILES["links"], "r", encoding="utf-8") as f:
            data = json.load(f)
        if not isinstance(data, list):
            return []
        out = [
            {
                "name": str(item.get("name", "")).strip(),
                "link": str(item.get("link", "")).strip(),
                "icon": str(item.get("icon", "")).strip(),
            }
            for item in data
            if isinstance(item, dict)
        ]
        return [item for item in out if any(item.values())]
    except (OSError, json.JSONDecodeError):
        return []


def save_links(links: List[Dict[str, str]]):
    os.makedirs(os.path.dirname(FILES["links"]), exist_ok=True)
    clean = [
        item
        for link_data in links
        if any((item := {
            "name": (link_data.get("name") or "").strip(),
            "link": (link_data.get("link") or "").strip(),
            "icon": (link_data.get("icon") or "").strip(),
        }).values())
    ]
    with open(FILES["links"], "w", encoding="utf-8") as f:
        json.dump(clean, f, ensure_ascii=False, indent=2)


def load_app_settings() -> Dict[str, bool]:
    settings = {"preview_before_save": True}
    try:
        with open(APP_SETTINGS_CONFIG, "r", encoding="utf-8") as stream:
            saved = json.load(stream)
        if isinstance(saved, dict) and isinstance(saved.get("preview_before_save"), bool):
            settings["preview_before_save"] = saved["preview_before_save"]
    except (OSError, json.JSONDecodeError):
        pass
    return settings


def save_app_settings(settings: Dict[str, bool]):
    os.makedirs(os.path.dirname(APP_SETTINGS_CONFIG), exist_ok=True)
    temporary_path = APP_SETTINGS_CONFIG + ".tmp"
    with open(temporary_path, "w", encoding="utf-8") as stream:
        json.dump(settings, stream, ensure_ascii=False, indent=2)
    os.replace(temporary_path, APP_SETTINGS_CONFIG)


# -----------------------------
# Windows acrylic helper (safe)
# -----------------------------
def _is_windows() -> bool:
    return os.name == "nt"


def _try_enable_acrylic(hwnd: int, dark: bool):
    if not _is_windows() or not hwnd:
        return
    try:
        user32 = ctypes.windll.user32

        class ACCENT_POLICY(ctypes.Structure):
            _fields_ = [
                ("AccentState", ctypes.c_int),
                ("AccentFlags", ctypes.c_int),
                ("GradientColor", ctypes.c_int),
                ("AnimationId", ctypes.c_int),
            ]

        class WINDOWCOMPOSITIONATTRIBDATA(ctypes.Structure):
            _fields_ = [
                ("Attribute", ctypes.c_int),
                ("Data", ctypes.c_void_p),
                ("SizeOfData", ctypes.c_size_t),
            ]

        WCA_ACCENT_POLICY = 19
        ACCENT_ENABLE_ACRYLICBLURBEHIND = 4

        if dark:
            r, g, b, a = 22, 24, 28, 210
        else:
            r, g, b, a = 255, 255, 255, 200

        gradient = (a << 24) | (b << 16) | (g << 8) | r

        accent = ACCENT_POLICY()
        accent.AccentState = ACCENT_ENABLE_ACRYLICBLURBEHIND
        accent.AccentFlags = 2
        accent.GradientColor = gradient
        accent.AnimationId = 0

        data = WINDOWCOMPOSITIONATTRIBDATA()
        data.Attribute = WCA_ACCENT_POLICY
        data.Data = ctypes.cast(ctypes.pointer(accent), ctypes.c_void_p)
        data.SizeOfData = ctypes.sizeof(accent)

        set_wca = getattr(user32, "SetWindowCompositionAttribute", None)
        if set_wca:
            set_wca(hwnd, ctypes.byref(data))
    except Exception:
        pass


# -----------------------------
# Dialogs
# -----------------------------
class ListEditorDialog(QDialog):
    def __init__(self, title: str, file_path: str, parent=None):
        super().__init__(parent)
        self.setWindowTitle(title)
        self.resize(720, 520)
        self.file_path = file_path

        self.listw = QListWidget()
        self.listw.setSelectionMode(QAbstractItemView.ExtendedSelection)

        self.btn_add = QPushButton("Добавить")
        self.btn_edit = QPushButton("Изменить")
        self.btn_del = QPushButton("Удалить")
        self.btn_up = QPushButton("Вверх")
        self.btn_down = QPushButton("Вниз")
        self.btn_save = QPushButton("Сохранить")
        self.btn_close = QPushButton("Закрыть")

        left = QVBoxLayout()
        left.addWidget(self.listw)

        right = QVBoxLayout()
        right.setSpacing(8)
        right.addWidget(self.btn_add)
        right.addWidget(self.btn_edit)
        right.addWidget(self.btn_del)
        right.addSpacing(10)
        right.addWidget(self.btn_up)
        right.addWidget(self.btn_down)
        right.addStretch()
        right.addWidget(self.btn_save)
        right.addWidget(self.btn_close)

        root = QVBoxLayout(self)
        wrapper = QWidget()
        row = QHBoxLayout(wrapper)
        row.setContentsMargins(14, 14, 14, 14)
        row.setSpacing(12)
        row.addLayout(left, 1)
        row.addLayout(right)
        root.addWidget(wrapper)

        self.btn_add.clicked.connect(self.add_item)
        self.btn_edit.clicked.connect(self.edit_item)
        self.btn_del.clicked.connect(self.del_item)
        self.btn_up.clicked.connect(self.move_up)
        self.btn_down.clicked.connect(self.move_down)
        self.btn_save.clicked.connect(self.save)
        self.btn_close.clicked.connect(self.close)

        self.load()

    def load(self):
        self.listw.clear()
        for x in read_lines(self.file_path):
            self.listw.addItem(QListWidgetItem(x))

    def add_item(self):
        text, ok = QInputDialog.getText(self, "Добавить", "Введите пункт:")
        if ok and text.strip():
            self.listw.addItem(QListWidgetItem(text.strip()))

    def edit_item(self):
        items = self.listw.selectedItems()
        if len(items) != 1:
            QMessageBox.information(self, "Изменить", "Выберите ровно один пункт.")
            return
        cur = items[0].text()
        text, ok = QInputDialog.getText(self, "Изменить", "Измените пункт:", text=cur)
        if ok and text.strip():
            items[0].setText(text.strip())

    def del_item(self):
        for it in self.listw.selectedItems():
            row = self.listw.row(it)
            self.listw.takeItem(row)

    def move_up(self):
        row = self.listw.currentRow()
        if row <= 0:
            return
        it = self.listw.takeItem(row)
        self.listw.insertItem(row - 1, it)
        self.listw.setCurrentRow(row - 1)

    def move_down(self):
        row = self.listw.currentRow()
        if row < 0 or row >= self.listw.count() - 1:
            return
        it = self.listw.takeItem(row)
        self.listw.insertItem(row + 1, it)
        self.listw.setCurrentRow(row + 1)

    def save(self):
        items = [self.listw.item(i).text().strip() for i in range(self.listw.count())]
        items = [x for x in items if x]
        write_lines(self.file_path, items)
        QMessageBox.information(self, "Сохранено", "Список сохранён.")


class LinksEditorDialog(QDialog):
    def __init__(self, links: List[Dict[str, str]], parent=None):
        super().__init__(parent)
        self.setWindowTitle("Редактор ссылок")
        self.resize(860, 580)
        self.links = [dict(x) for x in links]

        self.listw = QListWidget()
        self.listw.setSelectionMode(QAbstractItemView.SingleSelection)

        self.btn_add = QPushButton("Добавить")
        self.btn_edit = QPushButton("Изменить")
        self.btn_icon = QPushButton("Иконка…")
        self.btn_clear_icon = QPushButton("Сбросить иконку")
        self.btn_del = QPushButton("Удалить")
        self.btn_up = QPushButton("Вверх")
        self.btn_down = QPushButton("Вниз")
        self.btn_save = QPushButton("Сохранить")
        self.btn_close = QPushButton("Закрыть")

        left = QVBoxLayout()
        left.addWidget(self.listw)

        right = QVBoxLayout()
        right.setSpacing(8)
        right.addWidget(self.btn_add)
        right.addWidget(self.btn_edit)
        right.addWidget(self.btn_icon)
        right.addWidget(self.btn_clear_icon)
        right.addWidget(self.btn_del)
        right.addSpacing(10)
        right.addWidget(self.btn_up)
        right.addWidget(self.btn_down)
        right.addStretch()
        right.addWidget(self.btn_save)
        right.addWidget(self.btn_close)

        wrapper = QWidget()
        row = QHBoxLayout(wrapper)
        row.setContentsMargins(14, 14, 14, 14)
        row.setSpacing(12)
        row.addLayout(left, 1)
        row.addLayout(right)

        root = QVBoxLayout(self)
        root.addWidget(wrapper)

        self.btn_add.clicked.connect(self.add_item)
        self.btn_edit.clicked.connect(self.edit_item)
        self.btn_icon.clicked.connect(self.change_icon_for_selected)
        self.btn_clear_icon.clicked.connect(self.clear_icon_for_selected)
        self.btn_del.clicked.connect(self.del_item)
        self.btn_up.clicked.connect(self.move_up)
        self.btn_down.clicked.connect(self.move_down)
        self.btn_save.clicked.connect(self.accept)
        self.btn_close.clicked.connect(self.reject)

        self.refresh()

    def _pick_icon(self, current: str = "") -> str:
        current_path = resolve_portable_path(current)
        start_dir = os.path.dirname(current_path) if current_path and os.path.exists(current_path) else ICONS_DIR
        path, _ = QFileDialog.getOpenFileName(
            self,
            "Выберите иконку",
            start_dir,
            "Images (*.png *.ico *.jpg *.jpeg *.bmp *.webp);;All files (*.*)"
        )
        return portable_stored_path(path)

    def refresh(self):
        self.listw.clear()
        for t in self.links:
            name = (t.get("name") or "").strip() or "(без названия)"
            link = (t.get("link") or "").strip()
            it = QListWidgetItem(f"{name}  —  {link}")
            ico = icon_or_empty((t.get("icon") or "").strip())
            if not ico.isNull():
                it.setIcon(ico)
            self.listw.addItem(it)

    def add_item(self):
        name, ok = QInputDialog.getText(self, "Название", "Название ссылки:")
        if not ok:
            return
        link, ok2 = QInputDialog.getText(self, "Ссылка/путь", "Ссылка или путь к файлу:")
        if not ok2:
            return
        icon = self._pick_icon("")
        self.links.append({"name": name.strip(), "link": link.strip(), "icon": icon})
        self.refresh()

    def edit_item(self):
        row = self.listw.currentRow()
        if row < 0:
            QMessageBox.information(self, "Изменить", "Выберите ссылку.")
            return

        cur = self.links[row]
        name, ok = QInputDialog.getText(self, "Название", "Название ссылки:", text=cur.get("name", ""))
        if not ok:
            return
        link, ok2 = QInputDialog.getText(self, "Ссылка/путь", "Ссылка или путь к файлу:", text=cur.get("link", ""))
        if not ok2:
            return

        self.links[row]["name"] = name.strip()
        self.links[row]["link"] = link.strip()
        self.refresh()
        self.listw.setCurrentRow(row)

    def change_icon_for_selected(self):
        row = self.listw.currentRow()
        if row < 0:
            QMessageBox.information(self, "Иконка", "Выберите ссылку.")
            return
        cur = self.links[row]
        picked = self._pick_icon(cur.get("icon", ""))
        if picked:
            self.links[row]["icon"] = picked
            self.refresh()
            self.listw.setCurrentRow(row)

    def clear_icon_for_selected(self):
        row = self.listw.currentRow()
        if row < 0:
            return
        self.links[row]["icon"] = ""
        self.refresh()
        self.listw.setCurrentRow(row)

    def del_item(self):
        row = self.listw.currentRow()
        if row < 0:
            return
        self.links.pop(row)
        self.refresh()

    def move_up(self):
        row = self.listw.currentRow()
        if row <= 0:
            return
        self.links[row - 1], self.links[row] = self.links[row], self.links[row - 1]
        self.refresh()
        self.listw.setCurrentRow(row - 1)

    def move_down(self):
        row = self.listw.currentRow()
        if row < 0 or row >= len(self.links) - 1:
            return
        self.links[row + 1], self.links[row] = self.links[row], self.links[row + 1]
        self.refresh()
        self.listw.setCurrentRow(row + 1)


# -----------------------------
# Search combo
# -----------------------------
class SearchCombo(QComboBox):
    def __init__(self, items: List[str], editor_title: str, editor_file: str, parent=None):
        super().__init__(parent)
        self.setEditable(True)
        self.setInsertPolicy(QComboBox.NoInsert)
        self.setMaxVisibleItems(14)

        self._editor_title = editor_title
        self._editor_file = editor_file

        self._all_items: List[str] = []
        self._norm_map: List[tuple[str, str]] = []

        self._hint_model = QStringListModel(self)
        self._completer = QCompleter(self._hint_model, self)
        self._completer.setCompletionMode(QCompleter.PopupCompletion)
        self._completer.setCaseSensitivity(Qt.CaseInsensitive)
        self._completer.setFilterMode(Qt.MatchContains)

        self._completer.activated.connect(self._on_completer_activated)
        self.lineEdit().setCompleter(self._completer)

        self.lineEdit().textEdited.connect(self._on_text_edited)
        self.lineEdit().returnPressed.connect(self._on_return_pressed)

        gear_icon = self.style().standardIcon(QStyle.SP_FileDialogDetailedView)
        act = QAction(gear_icon, "Редактировать список", self)
        act.setToolTip("Редактировать список")
        act.setIconVisibleInMenu(False)
        act.triggered.connect(self.open_editor)
        self.lineEdit().addAction(act, QLineEdit.TrailingPosition)
        self.lineEdit().setClearButtonEnabled(False)

        self.set_items(items)

    def set_items(self, items: List[str]):
        cur = (self.currentText() or "").strip()
        self._all_items = list(dict.fromkeys(
            text for item in (items or []) if item and (text := item.strip())
        ))
        self._norm_map = [(normalize_search(x), x) for x in self._all_items]

        with QSignalBlocker(self):
            self.clear()
            self.addItems(self._all_items)

        self._hint_model.setStringList(self._all_items)
        self.setEditText(cur)

    def open_editor(self):
        dlg = ListEditorDialog(self._editor_title, self._editor_file, self.window())
        dlg.exec()
        self.set_items(read_lines(self._editor_file))
        self._show_hints()

    def _build_hints(self, typed: str) -> List[str]:
        qn = normalize_search(typed)
        if not qn:
            return self._all_items[:]
        hits = [orig for n, orig in self._norm_map if qn in n]
        if not hits:
            parts = [normalize_search(p) for p in re.split(r"\s+", typed.strip()) if p.strip()]
            parts = [p for p in parts if p]
            if parts:
                hits = [orig for n, orig in self._norm_map if all(p in n for p in parts)]
        return hits

    def _on_text_edited(self, text: str):
        self._hint_model.setStringList(self._build_hints(text))
        self._show_hints()

    def _show_hints(self):
        if self._hint_model.rowCount() > 0:
            self._completer.complete()

    def _on_completer_activated(self, text: str):
        text = (text or "").strip()
        if not text:
            return
        self.setEditText(text)
        idx = self.findText(text, Qt.MatchExactly)
        if idx >= 0:
            self.setCurrentIndex(idx)

    def _on_return_pressed(self):
        typed = (self.lineEdit().text() or "").strip()
        if not typed:
            return
        hints = self._build_hints(typed)
        if hints:
            self._on_completer_activated(hints[0])
            return
        idx = self.findText(typed, Qt.MatchExactly)
        if idx >= 0:
            self.setCurrentIndex(idx)
        else:
            self.setEditText(typed)

    def text(self) -> str:
        return (self.currentText() or "").strip()


class PreviewRenderThread(QThread):
    """Render a DOCX without blocking construction of the preview dialog."""

    def __init__(self, temporary_path: str, parent=None):
        super().__init__(parent)
        self.temporary_path = temporary_path
        self.pdf_path = None
        self.error = None

    def run(self):
        try:
            self.pdf_path = render_docx_to_pdf(self.temporary_path)
        except Exception as error:
            self.error = error


class ActPreviewDialog(QDialog):
    """Show the populated document as real rendered pages before saving."""

    def __init__(self, temporary_path: str, parent=None):
        super().__init__(parent)
        self.temporary_path = temporary_path
        self.pdf_path = None
        self.pdf_bytes = None
        self.pdf_buffer = None
        self.pdf_document = None
        self._loading_step = 0
        self._cancel_requested = False
        self.setWindowTitle("Предпросмотр акта")
        self.resize(980, 780)

        layout = QVBoxLayout(self)
        info = QLabel(
            "Проверьте полностью заполненный акт. Можно открыть временный DOCX, "
            "вернуться к редактированию или подтвердить окончательное сохранение."
        )
        info.setWordWrap(True)
        layout.addWidget(info)

        self.loading_label = QLabel("Подготавливаем визуальный предпросмотр…")
        self.loading_label.setAlignment(Qt.AlignCenter)
        layout.addWidget(self.loading_label, 1)
        self.preview = QPdfView(self)
        self.preview.setPageMode(QPdfView.PageMode.MultiPage)
        self.preview.setZoomMode(QPdfView.ZoomMode.FitToWidth)
        self.preview.hide()
        layout.addWidget(self.preview, 1)

        controls = QHBoxLayout()
        self.open_button = QPushButton("Открыть документ")
        self.open_button.setToolTip("Станет доступно после подготовки предпросмотра")
        self.open_button.clicked.connect(lambda: open_file(self.temporary_path))
        controls.addWidget(self.open_button)

        self.fit_button = QPushButton("По ширине")
        self.fit_button.clicked.connect(
            lambda: self.preview.setZoomMode(QPdfView.ZoomMode.FitToWidth)
        )
        self.zoom_out_button = QPushButton("−")
        self.zoom_out_button.setFixedWidth(42)
        self.zoom_out_button.clicked.connect(lambda: self._zoom(0.85))
        self.zoom_in_button = QPushButton("+")
        self.zoom_in_button.setFixedWidth(42)
        self.zoom_in_button.clicked.connect(lambda: self._zoom(1.15))
        controls.addWidget(self.fit_button)
        controls.addWidget(self.zoom_out_button)
        controls.addWidget(self.zoom_in_button)
        controls.addStretch(1)

        self.back_button = QPushButton("Вернуться к редактированию")
        self.back_button.clicked.connect(self.reject)
        self.save_button = QPushButton("Подтвердить и сохранить")
        self.save_button.setDefault(True)
        self.save_button.clicked.connect(self.accept)
        controls.addWidget(self.back_button)
        controls.addWidget(self.save_button)
        layout.addLayout(controls)

        for button in (
            self.open_button,
            self.fit_button,
            self.zoom_out_button,
            self.zoom_in_button,
            self.save_button,
        ):
            button.setEnabled(False)

        self.loading_timer = QTimer(self)
        self.loading_timer.setInterval(350)
        self.loading_timer.timeout.connect(self._animate_loading)
        self.loading_timer.start()

        self.render_thread = PreviewRenderThread(temporary_path, self)
        self.render_thread.finished.connect(self._render_finished)
        self.render_thread.start()

    def _animate_loading(self):
        self._loading_step = (self._loading_step + 1) % 4
        self.loading_label.setText(
            "Подготавливаем визуальный предпросмотр" + "." * self._loading_step
        )

    def _render_finished(self):
        self.loading_timer.stop()
        if self._cancel_requested:
            super().done(QDialog.Rejected)
            return
        if self.render_thread.error is not None:
            QMessageBox.critical(
                self,
                "Ошибка предпросмотра",
                str(self.render_thread.error),
            )
            super().done(QDialog.Rejected)
            return
        try:
            self._load_pdf(self.render_thread.pdf_path)
        except (OSError, PreviewRenderError) as error:
            QMessageBox.critical(self, "Ошибка предпросмотра", str(error))
            super().done(QDialog.Rejected)

    def _load_pdf(self, pdf_path: str):
        self.pdf_path = pdf_path
        with open(self.pdf_path, "rb") as pdf_file:
            self.pdf_bytes = QByteArray(pdf_file.read())
        self.pdf_buffer = QBuffer(self)
        self.pdf_buffer.setData(self.pdf_bytes)
        if not self.pdf_buffer.open(QIODevice.ReadOnly):
            raise PreviewRenderError("Не удалось открыть PDF предпросмотра в памяти.")

        self.pdf_document = QPdfDocument(self)
        load_error = self.pdf_document.load(self.pdf_buffer)
        if load_error is not None and load_error != QPdfDocument.Error.None_:
            raise PreviewRenderError(f"Qt не смог открыть PDF предпросмотра: {load_error}")

        self.preview.setDocument(self.pdf_document)
        self.loading_label.hide()
        self.preview.show()
        self.open_button.setToolTip("")
        for button in (
            self.open_button,
            self.fit_button,
            self.zoom_out_button,
            self.zoom_in_button,
            self.save_button,
        ):
            button.setEnabled(True)

    def _zoom(self, factor: float):
        self.preview.setZoomMode(QPdfView.ZoomMode.Custom)
        self.preview.setZoomFactor(max(0.25, min(4.0, self.preview.zoomFactor() * factor)))

    def done(self, result: int):
        if self.render_thread.isRunning():
            self._cancel_requested = True
            self.loading_label.setText("Завершаем подготовку предпросмотра…")
            self.save_button.setEnabled(False)
            self.back_button.setEnabled(False)
            return
        if self.pdf_document is not None:
            self.preview.setDocument(None)
            self.pdf_document.close()
        if self.pdf_buffer is not None:
            self.pdf_buffer.close()
        super().done(result)


# -----------------------------
# Main window
# -----------------------------
class MainWindow(QMainWindow):
    def __init__(self):
        super().__init__()
        ensure_dirs()
        ensure_default_icons()
        ensure_arrow_icons()

        self.setWindowTitle("ActGeneratorPC")
        self.resize(1100, 820)

        self.dark_theme = True
        self.links = load_links()
        self.app_settings = load_app_settings()

        self.data_lists = {
            key: read_lines(path)
            for key, path in FILES.items()
            if key != "links"
        }

        self._menus_for_mask = set()

        self._build_ui()
        self.apply_theme(True)
        self.refresh_links_menu()

    # ---- Menu round mask helpers ----
    def _apply_menu_round_mask(self, menu: QMenu):
        try:
            r = 14
            rect = menu.rect()
            if rect.width() <= 2 or rect.height() <= 2:
                return
            path = QPainterPath()
            path.addRoundedRect(rect, r, r)
            menu.setMask(QRegion(path.toFillPolygon().toPolygon()))
        except Exception:
            pass

    def eventFilter(self, obj, event):
        if isinstance(obj, QMenu) and obj in getattr(self, "_menus_for_mask", set()):
            if event.type() in (QEvent.Show, QEvent.Resize):
                QTimer.singleShot(0, lambda m=obj: self._apply_menu_round_mask(m))
        return super().eventFilter(obj, event)

    def _beautify_topbar(self):
        eff = QGraphicsDropShadowEffect(self.topbar)
        eff.setBlurRadius(26)
        eff.setOffset(0, 10)
        eff.setColor(QColor(0, 0, 0, 110))
        self.topbar.setGraphicsEffect(eff)

    def _beautify_menu(self, menu: QMenu):
        menu.setAttribute(Qt.WA_TranslucentBackground, True)
        menu.setAutoFillBackground(False)

        eff = QGraphicsDropShadowEffect(menu)
        eff.setBlurRadius(28)
        eff.setOffset(0, 10)
        eff.setColor(QColor(0, 0, 0, 140))
        menu.setGraphicsEffect(eff)

        def _apply_acrylic_now():
            try:
                hwnd = int(menu.winId())
                _try_enable_acrylic(hwnd, self.dark_theme)
            except Exception:
                pass
            QTimer.singleShot(0, lambda m=menu: self._apply_menu_round_mask(m))

        menu.aboutToShow.connect(_apply_acrylic_now)
        menu.installEventFilter(self)
        self._menus_for_mask.add(menu)

    def _menu_css(self) -> str:
        if self.dark_theme:
            menu_bg = "rgba(22, 24, 28, 210)"
            menu_border = "rgba(255, 255, 255, 35)"
            menu_sep = "rgba(255, 255, 255, 28)"
            item_hover = "rgba(255, 255, 255, 18)"
            item_press = "rgba(45, 127, 249, 55)"
            text = "#e8eaed"
            disabled = "rgba(232, 234, 237, 120)"
        else:
            menu_bg = "rgba(255, 255, 255, 220)"
            menu_border = "rgba(0, 0, 0, 38)"
            menu_sep = "rgba(0, 0, 0, 26)"
            item_hover = "rgba(0, 0, 0, 6)"
            item_press = "rgba(45, 127, 249, 45)"
            text = "#111111"
            disabled = "rgba(17, 17, 17, 120)"

        return f"""
            QMenu {{
                background: {menu_bg};
                color: {text};
                border: 1px solid {menu_border};
                border-radius: 14px;
                padding: 8px;
                background-clip: padding;
            }}
            QMenu::icon {{
                width: 18px;
                height: 18px;
                padding-left: 2px;
            }}
            QMenu::item {{
                padding: 10px 14px;
                padding-left: 34px;
                border-radius: 10px;
                margin: 2px 2px;
                spacing: 10px;
            }}
            QMenu::item:selected {{
                background: {item_hover};
            }}
            QMenu::item:pressed {{
                background: {item_press};
            }}
            QMenu::item:disabled {{
                color: {disabled};
            }}
            QMenu::separator {{
                height: 1px;
                background: {menu_sep};
                margin: 6px 6px;
            }}
        """

    # ---------------- UI ----------------
    def check_for_updates(self):
        manager = getattr(self, "update_manager", None)
        if manager is None:
            QMessageBox.warning(
                self,
                "Проверка обновлений",
                "Модуль обновления ещё не запущен. Повторите попытку через несколько секунд.",
            )
            return
        manager.check_async(manual=True)

    def choose_excel_workbook(self) -> str:
        current = load_excel_path(EXCEL_EXPORT_CONFIG)
        start_at = current if current and os.path.exists(current) else APP_DIR
        path, _ = QFileDialog.getOpenFileName(
            self,
            "Выберите Excel-таблицу для автозаполнения",
            start_at,
            "Excel (*.xlsx)",
        )
        if not path:
            return ""
        try:
            save_excel_path(EXCEL_EXPORT_CONFIG, path)
        except OSError as exc:
            QMessageBox.critical(
                self,
                "Настройка Excel",
                f"Не удалось сохранить путь к таблице:\n{exc}",
            )
            return ""
        self.lbl_status.setText(f"Excel-таблица автозаполнения: {path}")
        return path

    def use_embedded_excel_workbook(self) -> str:
        try:
            path = ensure_embedded_excel_workbook(
                EXCEL_TEMPLATE, EMBEDDED_EXCEL_WORKBOOK
            )
            save_excel_path(EXCEL_EXPORT_CONFIG, "")
        except (ExcelExportError, OSError) as exc:
            QMessageBox.critical(
                self,
                "Встроенная Excel-таблица",
                f"Не удалось включить встроенную таблицу:\n{exc}",
            )
            return ""
        self.lbl_status.setText(f"Используется встроенная Excel-таблица: {path}")
        return path

    def active_excel_workbook(self) -> str:
        return resolve_excel_workbook(
            EXCEL_EXPORT_CONFIG, EXCEL_TEMPLATE, EMBEDDED_EXCEL_WORKBOOK
        )

    def open_excel_workbook(self):
        try:
            path = self.active_excel_workbook()
        except ExcelExportError as exc:
            QMessageBox.critical(self, "Excel-таблица", str(exc))
            return
        open_file(path)

    def set_preview_enabled(self, enabled: bool):
        self.app_settings["preview_before_save"] = bool(enabled)
        try:
            save_app_settings(self.app_settings)
        except OSError as error:
            QMessageBox.warning(
                self,
                "Настройки",
                f"Не удалось сохранить настройку предпросмотра:\n{error}",
            )

    def _build_topbar(self) -> QWidget:
        bar = QFrame()
        bar.setObjectName("TopBar")
        lay = QHBoxLayout(bar)
        lay.setContentsMargins(10, 8, 10, 8)
        lay.setSpacing(8)

        # --- File menu
        self.menu_file = QMenu(self)
        self.act_check_updates = QAction("Проверить обновления", self)
        self.act_check_updates.triggered.connect(self.check_for_updates)
        self.menu_file.addAction(self.act_check_updates)
        self.menu_file.addSeparator()

        self.act_preview_enabled = QAction("Предпросмотр перед сохранением", self)
        self.act_preview_enabled.setCheckable(True)
        self.act_preview_enabled.setChecked(
            self.app_settings.get("preview_before_save", True)
        )
        self.act_preview_enabled.toggled.connect(self.set_preview_enabled)
        self.menu_file.addAction(self.act_preview_enabled)
        self.menu_file.addSeparator()

        self.act_select_excel = QAction("Выбрать другую Excel-таблицу…", self)
        self.act_select_excel.triggered.connect(self.choose_excel_workbook)
        self.menu_file.addAction(self.act_select_excel)
        self.act_use_embedded_excel = QAction(
            "Использовать встроенную Excel-таблицу", self
        )
        self.act_use_embedded_excel.triggered.connect(
            self.use_embedded_excel_workbook
        )
        self.menu_file.addAction(self.act_use_embedded_excel)
        self.act_open_excel = QAction("Открыть Excel-таблицу", self)
        self.act_open_excel.triggered.connect(self.open_excel_workbook)
        self.menu_file.addAction(self.act_open_excel)
        self.menu_file.addSeparator()

        act_exit = QAction("Выход", self)
        act_exit.triggered.connect(self.close)
        self.menu_file.addAction(act_exit)

        self.btn_file = QToolButton()
        self.btn_file.setText("Файл")
        self.btn_file.setIcon(self.style().standardIcon(QStyle.SP_DirHomeIcon))
        self.btn_file.setToolButtonStyle(Qt.ToolButtonTextBesideIcon)
        self.btn_file.setPopupMode(QToolButton.InstantPopup)
        self.btn_file.setMenu(self.menu_file)

        # --- Links menu (renamed from tables)
        self.menu_links = QMenu(self)
        self.act_edit_links = QAction("Добавить / редактировать ссылки", self)
        self.act_edit_links.triggered.connect(self.edit_links)
        self.menu_links.addAction(self.act_edit_links)
        self.menu_links.addSeparator()

        self.btn_links = QToolButton()
        self.btn_links.setText("Ссылки")
        self.btn_links.setIcon(icon_or_empty(ICON_EXCEL))
        self.btn_links.setToolButtonStyle(Qt.ToolButtonTextBesideIcon)
        self.btn_links.setPopupMode(QToolButton.InstantPopup)
        self.btn_links.setMenu(self.menu_links)

        # --- Lists menu
        self.menu_lists = QMenu(self)

        def add_list(title: str, key: str, icon_path: str):
            a = QAction(icon_or_empty(icon_path), f"Редактировать: {title}", self)
            a.setIconVisibleInMenu(True)
            a.triggered.connect(lambda: self.open_list_editor(title, key))
            self.menu_lists.addAction(a)

        add_list("Исполнители", "executors", ICON_USER)
        add_list("Станции", "locations", ICON_STATION)
        add_list("Типы оборудования", "models", ICON_EQUIPMENT)
        add_list("Выполненные работы", "work", ICON_SERVICES)
        add_list("Несоответствие", "issues", ICON_ISSUES)
        add_list("Проделанная работа", "done_work", ICON_DONE_WORK)
        add_list("Расходные материалы", "materials", ICON_MATERIALS)

        self.btn_lists = QToolButton()
        self.btn_lists.setText("Списки")
        self.btn_lists.setIcon(icon_or_empty(ICON_LISTS))
        self.btn_lists.setToolButtonStyle(Qt.ToolButtonTextBesideIcon)
        self.btn_lists.setPopupMode(QToolButton.InstantPopup)
        self.btn_lists.setMenu(self.menu_lists)

        # --- Theme button
        self.btn_theme = QToolButton()
        self.btn_theme.setText("Тема")
        self.btn_theme.setIcon(self.style().standardIcon(QStyle.SP_TitleBarShadeButton))
        self.btn_theme.setToolButtonStyle(Qt.ToolButtonTextBesideIcon)
        self.btn_theme.setCheckable(True)
        self.btn_theme.setChecked(True)
        self.btn_theme.clicked.connect(self._toggle_theme_from_button)

        for b in (self.btn_file, self.btn_links, self.btn_lists, self.btn_theme):
            b.setMinimumHeight(40)

        self.btn_file.setMinimumWidth(120)
        self.btn_links.setMinimumWidth(120)
        self.btn_lists.setMinimumWidth(120)
        self.btn_theme.setMinimumWidth(110)

        lay.addWidget(self.btn_file)
        lay.addWidget(self.btn_links)
        lay.addWidget(self.btn_lists)
        lay.addStretch(1)
        lay.addWidget(self.btn_theme)

        self._beautify_menu(self.menu_file)
        self._beautify_menu(self.menu_links)
        self._beautify_menu(self.menu_lists)

        return bar

    def _build_ui(self):
        central = QWidget()
        self.setCentralWidget(central)

        root = QVBoxLayout(central)
        root.setContentsMargins(16, 16, 16, 16)
        root.setSpacing(12)

        self.topbar = self._build_topbar()
        root.addWidget(self.topbar)
        self._beautify_topbar()

        form = QFormLayout()
        form.setHorizontalSpacing(14)
        form.setVerticalSpacing(12)
        root.addLayout(form, 1)

        self.ed_date = QLineEdit()
        self.ed_date.setPlaceholderText("Дата (например 04.02.2026)")
        self.ed_date.setText(datetime.date.today().strftime("%d.%m.%Y"))
        form.addRow(make_form_label("Дата:", ICON_CALENDAR), self.ed_date)

        self.cb_executor = SearchCombo(self.data_lists["executors"], "Исполнители", FILES["executors"])
        form.addRow(make_form_label("Исполнитель:", ICON_USER), self.cb_executor)

        self.cb_location = SearchCombo(self.data_lists["locations"], "Станции", FILES["locations"])
        form.addRow(make_form_label("Станция:", ICON_STATION), self.cb_location)

        self.cb_model = SearchCombo(self.data_lists["models"], "Типы оборудования", FILES["models"])
        self.cb_model.lineEdit().textChanged.connect(lambda _: self.update_template_hint())
        form.addRow(make_form_label("Тип оборудования:", ICON_EQUIPMENT), self.cb_model)

        self.ed_serial = QLineEdit()
        self.ed_serial.setPlaceholderText("№ ККТ")
        form.addRow(make_form_label("№ ККТ:", ICON_SERIAL), self.ed_serial)

        self.cb_work = SearchCombo(self.data_lists["work"], "Выполненные работы", FILES["work"])
        form.addRow(make_form_label("Выполненные работы:", ICON_SERVICES), self.cb_work)

        self.cb_issues = SearchCombo(self.data_lists["issues"], "Несоответствие", FILES["issues"])
        form.addRow(make_form_label("Несоответствие:", ICON_ISSUES), self.cb_issues)

        self.cb_done_work = SearchCombo(self.data_lists["done_work"], "Проделанная работа", FILES["done_work"])
        form.addRow(make_form_label("Проделанная работа:", ICON_DONE_WORK), self.cb_done_work)

        self.cb_materials = SearchCombo(self.data_lists["materials"], "Расходные материалы", FILES["materials"])
        form.addRow(make_form_label("Расходные материалы:", ICON_MATERIALS), self.cb_materials)

        self.sp_qty = QSpinBox()
        self.sp_qty.setRange(0, 999999)
        self.sp_qty.setValue(0)
        form.addRow(make_form_label("Количество (цифрами):", ICON_QTY), self.sp_qty)

        self.cb_template = QComboBox()
        self.cb_template.addItem("Авто", None)
        self.cb_template.addItem(
            f"АБП / МКТФ ({os.path.basename(TEMPLATE_TYPE1)})",
            TEMPLATE_TYPE1,
        )
        self.cb_template.addItem(
            f"Валидатор / MID ({os.path.basename(TEMPLATE_TYPE2)})",
            TEMPLATE_TYPE2,
        )
        self.cb_template.currentIndexChanged.connect(lambda _: self.update_template_hint())
        form.addRow(make_form_label("Выбор шаблона:", ICON_DOCX), self.cb_template)
        self.update_template_hint()

        self.btn_create = QPushButton("Создать акт")
        self.btn_create.setMinimumHeight(44)
        self.btn_create.clicked.connect(self.create_act)
        root.addWidget(self.btn_create)

        self.btn_open_output = QPushButton("Открыть папку с актами")
        self.btn_open_output.setMinimumHeight(40)
        self.btn_open_output.clicked.connect(lambda: open_folder(OUTPUT_DIR))
        root.addWidget(self.btn_open_output)

        self.lbl_status = QLabel("")
        self.lbl_status.setWordWrap(True)
        root.addWidget(self.lbl_status)

    # ---------------- Theme / Styles ----------------
    def _toggle_theme_from_button(self):
        self.dark_theme = self.btn_theme.isChecked()
        self.apply_theme(self.dark_theme)

    def apply_theme(self, dark: bool):
        QApplication.instance().setStyle("Fusion")

        pal = QPalette()
        if not dark:
            pal.setColor(QPalette.Window, QColor("#f5f6f7"))
            pal.setColor(QPalette.WindowText, QColor("#111111"))
            pal.setColor(QPalette.Base, QColor("#ffffff"))
            pal.setColor(QPalette.Text, QColor("#111111"))
            pal.setColor(QPalette.Button, QColor("#ffffff"))
            pal.setColor(QPalette.ButtonText, QColor("#111111"))
            pal.setColor(QPalette.Highlight, QColor("#2d7ff9"))
            pal.setColor(QPalette.HighlightedText, QColor("#ffffff"))
            border = "#cfd3d7"
            bg = "#ffffff"
            dd_border = "#e3e6ea"
            popup_bg = "#ffffff"
            popup_fg = "#111111"
            spin_up = SPIN_UP_LIGHT_PNG
            spin_down = SPIN_DOWN_LIGHT_PNG
        else:
            pal.setColor(QPalette.Window, QColor("#121417"))
            pal.setColor(QPalette.WindowText, QColor("#e6e6e6"))
            pal.setColor(QPalette.Base, QColor("#1b1f24"))
            pal.setColor(QPalette.Text, QColor("#e6e6e6"))
            pal.setColor(QPalette.Button, QColor("#1b1f24"))
            pal.setColor(QPalette.ButtonText, QColor("#e6e6e6"))
            pal.setColor(QPalette.Highlight, QColor("#2d7ff9"))
            pal.setColor(QPalette.HighlightedText, QColor("#ffffff"))
            border = "#2b3138"
            bg = "#1b1f24"
            dd_border = "#2b3138"
            popup_bg = "#1b1f24"
            popup_fg = "#e6e6e6"
            spin_up = SPIN_UP_DARK_PNG
            spin_down = SPIN_DOWN_DARK_PNG

        if dark:
            tb_bg = "rgba(27, 31, 36, 210)"
            tb_border = "rgba(255, 255, 255, 22)"
            btn_bg = "rgba(255, 255, 255, 10)"
            btn_hover = "rgba(255, 255, 255, 18)"
            btn_press = "rgba(255, 255, 255, 26)"
            btn_checked = "rgba(45, 127, 249, 40)"
            btn_checked_border = "rgba(45, 127, 249, 120)"
        else:
            tb_bg = "rgba(255, 255, 255, 235)"
            tb_border = "rgba(0, 0, 0, 18)"
            btn_bg = "rgba(0, 0, 0, 4)"
            btn_hover = "rgba(0, 0, 0, 7)"
            btn_press = "rgba(0, 0, 0, 10)"
            btn_checked = "rgba(45, 127, 249, 22)"
            btn_checked_border = "rgba(45, 127, 249, 120)"

        QApplication.instance().setPalette(pal)

        arrow_png = ICON_ARROW_DOWN if os.path.exists(ICON_ARROW_DOWN) else ""
        arrow_png_qss = arrow_png.replace("\\", "/") if arrow_png else ""
        spin_up_qss = spin_up.replace("\\", "/")
        spin_down_qss = spin_down.replace("\\", "/")

        down_arrow_rule = f"""
        QComboBox::down-arrow {{
            image: url("{arrow_png_qss}");
            width: 18px;
            height: 18px;
            subcontrol-position: center;
            subcontrol-origin: content;
        }}
        """ if arrow_png_qss else ""

        self.setStyleSheet(f"""
            #TopBar {{
                background: {tb_bg};
                border: 1px solid {tb_border};
                border-radius: 16px;
                padding: 2px;
            }}

            QToolButton {{
                color: {popup_fg};
                background: {btn_bg};
                border: 1px solid rgba(255,255,255,0);
                padding: 8px 12px;
                border-radius: 12px;
                font-weight: 600;
            }}
            QToolButton:hover {{ background: {btn_hover}; }}
            QToolButton:pressed {{ background: {btn_press}; }}
            QToolButton:checked {{
                background: {btn_checked};
                border: 1px solid {btn_checked_border};
            }}
            QToolButton::menu-indicator {{
                image: none;
                width: 0px;
            }}

            {self._menu_css()}

            QComboBox {{
                padding: 7px 30px 7px 10px;
                border-radius: 10px;
                border: 1px solid {border};
                min-height: 34px;
                background: {bg};
            }}
            QComboBox::drop-down {{
                subcontrol-origin: padding;
                subcontrol-position: top right;
                width: 30px;
                border-left: 1px solid {dd_border};
            }}
            {down_arrow_rule}

            QComboBox QLineEdit {{
                padding-right: 8px;
                background: transparent;
                border: none;
            }}

            QLineEdit, QSpinBox {{
                padding: 7px;
                border-radius: 10px;
                border: 1px solid {border};
                min-height: 34px;
                background: {bg};
            }}

            QComboBox QAbstractItemView {{
                background: {popup_bg};
                color: {popup_fg};
                selection-background-color: #2d7ff9;
                selection-color: #ffffff;
            }}

            QSpinBox {{
                padding-right: 44px;
            }}
            QSpinBox::up-button {{
                subcontrol-origin: border;
                subcontrol-position: top right;
                width: 38px;
                border-left: 1px solid {dd_border};
                border-top-right-radius: 10px;
                background: {bg};
            }}
            QSpinBox::down-button {{
                subcontrol-origin: border;
                subcontrol-position: bottom right;
                width: 38px;
                border-left: 1px solid {dd_border};
                border-bottom-right-radius: 10px;
                background: {bg};
            }}
            QSpinBox::up-arrow {{
                image: url("{spin_up_qss}");
                width: 18px;
                height: 18px;
            }}
            QSpinBox::down-arrow {{
                image: url("{spin_down_qss}");
                width: 18px;
                height: 18px;
            }}

            QPushButton {{
                background: #2d7ff9;
                color: #ffffff;
                font-weight: 700;
                border: none;
                border-radius: 12px;
                padding: 10px 14px;
            }}
            QPushButton:hover {{ background: #276fe0; }}
        """)

    # ---------------- Menus data ----------------
    def refresh_links_menu(self):
        self.menu_links.clear()
        self.menu_links.addAction(self.act_edit_links)
        self.menu_links.addSeparator()

        self.links = load_links()
        if not self.links:
            a = QAction("(ссылки не добавлены)", self)
            a.setEnabled(False)
            self.menu_links.addAction(a)
            return

        default_icon = icon_or_empty(ICON_XLS) if os.path.exists(ICON_XLS) else QIcon()

        for t in self.links:
            name = (t.get("name") or "").strip() or "Ссылка"
            link = (t.get("link") or "").strip()

            ico_path = (t.get("icon") or "").strip()
            ico = icon_or_empty(ico_path) if ico_path else default_icon

            a = QAction(ico, name, self)
            a.setIconVisibleInMenu(True)
            a.triggered.connect(partial(open_any, link))
            self.menu_links.addAction(a)

    def open_list_editor(self, title: str, key: str):
        dlg = ListEditorDialog(title, FILES[key], self)
        dlg.exec()
        self.data_lists[key] = read_lines(FILES[key])
        combos = {
            "executors": self.cb_executor,
            "locations": self.cb_location,
            "models": self.cb_model,
            "work": self.cb_work,
            "issues": self.cb_issues,
            "done_work": self.cb_done_work,
            "materials": self.cb_materials,
        }
        combos[key].set_items(self.data_lists[key])
        if key == "models":
            self.update_template_hint()

    def edit_links(self):
        dlg = LinksEditorDialog(load_links(), self)
        if dlg.exec() == QDialog.Accepted:
            save_links(dlg.links)
            self.refresh_links_menu()
            QMessageBox.information(self, "Ссылки", "Ссылки сохранены.")

    # ---------------- Act generation ----------------
    def pick_template_for_equipment_type(self, equipment_type: str) -> str:
        """
        ✅ Один тип (TEMPLATE_TYPE1) для:
        МКТФ, АБП-БН, АБП-09, АБП-09-М3, АН-15, БПА-20-БН1
        """
        t = (equipment_type or "").upper().replace(" ", "")

        # Все варианты, которые должны идти в TEMPLATE_TYPE1
        type1_tokens = (
            "МКТФ", "MKTF",
            "АБП",          # покроет АБП-БН, АБП-09, АБП-09-М3 и т.п.
            "АН-15", "АН15",
            "БПА",          # покроет БПА-20-БН1
        )

        if any(tok in t for tok in type1_tokens):
            return TEMPLATE_TYPE1

        return TEMPLATE_TYPE2

    def selected_template(self, equipment_type: str) -> str:
        """Возвращает выбранный вручную шаблон или результат автовыбора."""
        manual_template = self.cb_template.currentData()
        if manual_template:
            return manual_template
        return self.pick_template_for_equipment_type(equipment_type)

    def update_template_hint(self):
        tpl = self.selected_template(self.cb_model.text())
        self.cb_template.setToolTip(f"Будет использован: {os.path.basename(tpl)}")

    def create_act(self):
        date = (self.ed_date.text() or "").strip()
        executor = self.cb_executor.text()
        location = self.cb_location.text()
        model = self.cb_model.text()
        serial = (self.ed_serial.text() or "").strip()

        work = self.cb_work.text()
        issues = self.cb_issues.text()
        done_work = self.cb_done_work.text()

        materials = self.cb_materials.text()
        qty_int = int(self.sp_qty.value())

        materials_qty = str(qty_int) if qty_int > 0 else ""
        unit = "шт." if qty_int > 0 else ""

        tpl = self.selected_template(model)
        if not os.path.exists(tpl):
            QMessageBox.critical(self, "Нет шаблона", f"Не найден шаблон:\n{tpl}\n\nПоложи его в папку Data\\Templates.")
            return

        mapping = {
            "{DATE}": date,
            "{EXECUTOR}": executor,
            "{LOCATION}": location,
            "{MODEL}": model,
            "{SERIAL}": serial,
            "{WORK}": work,
            "{ISSUES}": issues,
            "{DONE_WORK}": done_work,
            "{MATERIALS}": materials,
            "{MATERIALS_QTY}": materials_qty,
            "{UNIT}": unit,
        }

        try:
            doc = Document(tpl)
            replace_placeholders_docx(doc, mapping)

            ts = datetime.datetime.now().strftime("%Y-%m-%d_%H-%M-%S")
            base = safe_filename(f"Акт_{location or 'БезСтанции'}_{serial or 'БезККТ'}_{ts}")
            out_path = os.path.join(OUTPUT_DIR, base + ".docx")

            os.makedirs(OUTPUT_DIR, exist_ok=True)
            if self.app_settings.get("preview_before_save", True):
                with tempfile.TemporaryDirectory(
                    prefix="actgenerator-preview-", ignore_cleanup_errors=True
                ) as preview_dir:
                    preview_path = os.path.join(preview_dir, base + ".docx")
                    doc.save(preview_path)
                    preview = ActPreviewDialog(preview_path, self)
                    if preview.exec() != QDialog.Accepted:
                        self.lbl_status.setText(
                            "Сохранение отменено. Можно изменить данные и снова открыть предпросмотр."
                        )
                        return
                    shutil.copy2(preview_path, out_path)
            else:
                doc.save(out_path)

            excel_status = "Excel: таблица не обновлена."
            try:
                workbook_path = self.active_excel_workbook()
                excel_result = append_act_to_workbook(
                    workbook_path,
                        {
                            "date": date,
                            "executor": executor,
                            "model": model,
                            "serial": serial,
                            "location": location,
                            "work": work,
                            "issues": issues,
                            "materials": materials,
                            "done_work": done_work,
                            "qty": qty_int if qty_int > 0 else "",
                            "area": area_for_station(location),
                            "notes": note_for_materials(materials),
                        },
                )
                excel_status = (
                    f"Excel: лист «{excel_result.sheet_name}», строка {excel_result.row_number}."
                )
            except ExcelExportError as exc:
                excel_status = f"Excel не обновлён: {exc}"
                QMessageBox.warning(
                    self,
                    "Акт сохранён, Excel не обновлён",
                    f"DOCX создан:\n{out_path}\n\n{exc}",
                )

            self.lbl_status.setText(f"✅ Файл создан: {out_path}\n{excel_status}")

            msg = QMessageBox(self)
            msg.setWindowTitle("Готово")
            msg.setText("Акт создан.")
            msg.setInformativeText(f"{out_path}\n\n{excel_status}")
            btn_open = msg.addButton("Открыть", QMessageBox.AcceptRole)
            msg.addButton("OK", QMessageBox.RejectRole)
            msg.exec()
            if msg.clickedButton() == btn_open:
                open_file(out_path)

        except Exception as e:
            QMessageBox.critical(self, "Ошибка", f"Не удалось создать акт:\n{e}")


def main():
    app = QApplication(sys.argv)
    app.setApplicationName("ActGeneratorPC")
    app.setOrganizationName("ActGeneratorPC")
    app.setWindowIcon(icon_or_empty(ICON_APP))
    app.setStyle("Fusion")
    window = MainWindow()
    window.show()
    QTimer.singleShot(2500, confirm_healthy_startup)

    # Keep the updater alive for the lifetime of the main window.
    window.update_manager = UpdateManager(window)
    QTimer.singleShot(1500, window.update_manager.check_async)

    sys.exit(app.exec())


if __name__ == "__main__":
    main()
